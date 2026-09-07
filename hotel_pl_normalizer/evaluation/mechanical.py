"""Cheap integrity checks around one completed P&L normalization run.

These checks deliberately do not repeat the mapper's accounting, hierarchy, or
coverage validators. They establish that the evaluator inputs belong together
and that the final workbook survived writing with the columns Codex reviews.
"""

from __future__ import annotations

import json
import math
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from pypdf import PdfReader

from hotel_pl_normalizer.evaluation.evidence import (
    EvidenceIndex,
    EvidenceIndexError,
    build_evidence_index,
)
from hotel_pl_normalizer.evaluation.models import MechanicalCheck, MechanicalStatus
from hotel_pl_normalizer.mapping.coa import canonical_coa_ids
from hotel_pl_normalizer.output import (
    FIRST_ACCOUNT_ROW,
    FIRST_PERIOD_COL,
    HEADER_ROW,
    ID_COL,
    _output_number,
)

SUPPORTED_SOURCE_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".pdf"}
_DETAIL_LIMIT = 25


def run_mechanical_checks(
    source_path: str | Path,
    run_log_path: str | Path,
    mapped_workbook_path: str | Path | None = None,
    *,
    run_log: Mapping[str, Any] | None = None,
    evidence_index: EvidenceIndex | None = None,
) -> list[MechanicalCheck]:
    """Run only post-run artifact and evidence checks for one P&L."""

    source = Path(source_path)
    log_path = Path(run_log_path)
    mapped = Path(mapped_workbook_path) if mapped_workbook_path is not None else None

    source_check = _check_source_artifact(source)
    disk_log, run_log_check = _read_run_log_artifact(log_path, supplied=run_log)
    effective_log = disk_log if disk_log is not None else run_log
    integrity_check, _ = _check_run_log_integrity(
        effective_log,
        source,
        supplied_index=evidence_index,
    )
    accepted = _mapping_accepted(effective_log)
    mapped_check = _check_mapped_artifact(mapped, accepted=accepted)
    correlation_check = _check_output_log_correlation(
        mapped,
        effective_log,
        mapped_readable=mapped_check.status == MechanicalStatus.PASS,
        accepted=accepted,
    )
    return [
        source_check,
        run_log_check,
        integrity_check,
        mapped_check,
        correlation_check,
    ]


def _check_source_artifact(path: Path) -> MechanicalCheck:
    if not path.is_file():
        return _check(
            "source_artifact",
            MechanicalStatus.FAIL,
            "Source P&L is missing or is not a file.",
            [str(path)],
        )
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SOURCE_SUFFIXES:
        return _check(
            "source_artifact",
            MechanicalStatus.FAIL,
            f"Source P&L format {suffix or '(none)'} is not supported.",
            [str(path)],
        )

    try:
        if suffix in {".xlsx", ".xlsm"}:
            book = openpyxl.load_workbook(path, read_only=True, data_only=False)
            try:
                count = len(book.sheetnames)
                if count == 0:
                    raise ValueError("workbook contains no sheets")
            finally:
                book.close()
            description = f"{count} worksheet(s)"
        elif suffix == ".xls":
            book = xlrd.open_workbook(str(path), on_demand=True)
            try:
                count = book.nsheets
                if count == 0:
                    raise ValueError("workbook contains no sheets")
            finally:
                book.release_resources()
            description = f"{count} worksheet(s)"
        else:
            reader = PdfReader(str(path))
            count = len(reader.pages)
            if count == 0:
                raise ValueError("PDF contains no pages")
            description = f"{count} page(s)"
    except Exception as exc:  # readers expose several format-specific exceptions
        return _check(
            "source_artifact",
            MechanicalStatus.FAIL,
            "Source P&L could not be opened.",
            [str(path), f"{type(exc).__name__}: {exc}"],
        )

    return _check(
        "source_artifact",
        MechanicalStatus.PASS,
        "Source P&L exists and is readable.",
        [str(path), f"format={suffix}", description],
    )


def _read_run_log_artifact(
    path: Path,
    *,
    supplied: Mapping[str, Any] | None,
) -> tuple[Mapping[str, Any] | None, MechanicalCheck]:
    if not path.is_file():
        return None, _check(
            "run_log_artifact",
            MechanicalStatus.FAIL,
            "Run log is missing or is not a file.",
            [str(path)],
        )
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return None, _check(
            "run_log_artifact",
            MechanicalStatus.FAIL,
            "Run log could not be read as JSON.",
            [str(path), f"{type(exc).__name__}: {exc}"],
        )
    if not isinstance(payload, Mapping):
        return None, _check(
            "run_log_artifact",
            MechanicalStatus.FAIL,
            "Run log JSON must contain one object.",
            [str(path)],
        )
    if supplied is not None and dict(payload) != dict(supplied):
        return payload, _check(
            "run_log_artifact",
            MechanicalStatus.FAIL,
            "The supplied run-log object differs from the file on disk.",
            [str(path)],
        )
    return payload, _check(
        "run_log_artifact",
        MechanicalStatus.PASS,
        "Run log exists and is readable JSON.",
        [str(path)],
    )


def _check_run_log_integrity(
    run_log: Mapping[str, Any] | None,
    source_path: Path,
    *,
    supplied_index: EvidenceIndex | None,
) -> tuple[MechanicalCheck, EvidenceIndex | None]:
    if not isinstance(run_log, Mapping):
        return (
            _check(
                "run_log_integrity",
                MechanicalStatus.FAIL,
                "Run-log schema could not be checked because no object was loaded.",
            ),
            None,
        )

    issues: list[str] = []
    if run_log.get("log_version") != 5:
        issues.append(
            f"log_version={run_log.get('log_version')!r}; version 5 is required"
        )

    source = run_log.get("source")
    period_ids: list[str] = []
    if not isinstance(source, Mapping):
        issues.append("source must be an object")
    else:
        source_name = source.get("name")
        if not isinstance(source_name, str) or not source_name.strip():
            issues.append("source.name must be a non-empty string")
        elif source_name.casefold() != source_path.name.casefold():
            issues.append(
                f"source.name={source_name!r} does not match {source_path.name!r}"
            )
        periods = source.get("periods")
        if not _is_sequence(periods) or not periods:
            issues.append("source.periods must be a non-empty list")
        else:
            for position, period in enumerate(periods):
                if not isinstance(period, Mapping):
                    issues.append(f"source.periods[{position}] must be an object")
                    continue
                period_id = str(period.get("period_id") or "").strip()
                label = str(period.get("label") or "").strip()
                if not period_id:
                    issues.append(f"source.periods[{position}] has no period_id")
                else:
                    period_ids.append(period_id)
                if not label:
                    issues.append(f"source.periods[{position}] has no label")
            if len(period_ids) != len(set(period_ids)):
                issues.append("source.periods contains duplicate period_id values")

    outcome = run_log.get("outcome")
    if not isinstance(outcome, Mapping):
        issues.append("outcome must be an object")
    elif not isinstance(outcome.get("accepted"), bool):
        issues.append("outcome.accepted must be a boolean")

    values_by_period = run_log.get("values_by_period")
    if not isinstance(values_by_period, Mapping):
        issues.append("values_by_period must be an object")
    else:
        if period_ids and set(values_by_period) != set(period_ids):
            issues.append(
                "values_by_period keys do not match source.periods period_id values"
            )
        for period_id, values in values_by_period.items():
            if not isinstance(values, Mapping):
                issues.append(f"values_by_period[{period_id!r}] must be an object")

    manifest = run_log.get("feedback_manifest")
    if not isinstance(manifest, Mapping):
        issues.append("feedback_manifest must be an object")
    elif not _is_sequence(manifest.get("findings")):
        issues.append("feedback_manifest.findings must be a list")

    index: EvidenceIndex | None = None
    try:
        if supplied_index is None:
            index = build_evidence_index(run_log)
        else:
            _validate_supplied_index(run_log, supplied_index)
            index = supplied_index
    except (EvidenceIndexError, TypeError, ValueError) as exc:
        issues.append(str(exc))

    if issues:
        return (
            _check(
                "run_log_integrity",
                MechanicalStatus.FAIL,
                f"Run log failed {len(issues)} schema/input integrity check(s).",
                _limited(issues),
            ),
            index,
        )
    return (
        _check(
            "run_log_integrity",
            MechanicalStatus.PASS,
            "Run log is a valid version-5 evaluator input for this source P&L.",
            [
                f"periods={len(period_ids)}",
                f"evidence_rows={len(index.rows) if index is not None else 0}",
                f"accounts={len(index.accounts) if index is not None else 0}",
            ],
        ),
        index,
    )


def _check_mapped_artifact(
    path: Path | None,
    *,
    accepted: bool | None,
) -> MechanicalCheck:
    if path is None:
        if accepted is False:
            return _check(
                "mapped_workbook_artifact",
                MechanicalStatus.PASS,
                "No mapped workbook is expected because the mapping was not accepted.",
                ["outcome.accepted=false"],
            )
        return _check(
            "mapped_workbook_artifact",
            MechanicalStatus.FAIL,
            "An accepted run requires a mapped workbook for evaluation.",
        )
    if not path.is_file():
        return _check(
            "mapped_workbook_artifact",
            MechanicalStatus.FAIL,
            "Mapped workbook is missing or is not a file.",
            [str(path)],
        )
    if path.suffix.lower() != ".xlsx":
        return _check(
            "mapped_workbook_artifact",
            MechanicalStatus.FAIL,
            "Mapped output must be an .xlsx workbook.",
            [str(path)],
        )
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            missing_sheets = [
                sheet_name
                for sheet_name in ("COA", "KHP Model Accounts", "Run Notes")
                if sheet_name not in book.sheetnames
            ]
            if missing_sheets:
                raise ValueError(
                    f"missing required sheet(s): {', '.join(missing_sheets)}"
                )
            coa = book["COA"]
            expected_headers = {
                "B2": "COA_ID",
                "W2": "Mapped Labels",
                "X2": "MODEL FEEDBACK",
            }
            for coordinate, expected in expected_headers.items():
                actual = str(coa[coordinate].value or "")
                if expected.casefold() not in actual.casefold():
                    raise ValueError(
                        f"COA {coordinate} must contain a {expected!r} header"
                    )
        finally:
            book.close()
    except Exception as exc:  # openpyxl exposes ZIP, XML, and IO exceptions
        return _check(
            "mapped_workbook_artifact",
            MechanicalStatus.FAIL,
            "Mapped workbook could not be reopened as a standard output.",
            [str(path), f"{type(exc).__name__}: {exc}"],
        )
    return _check(
        "mapped_workbook_artifact",
        MechanicalStatus.PASS,
        "Mapped workbook exists and reopens successfully.",
        [str(path)],
    )


def _check_output_log_correlation(
    path: Path | None,
    run_log: Mapping[str, Any] | None,
    *,
    mapped_readable: bool,
    accepted: bool | None,
) -> MechanicalCheck:
    if path is None and accepted is False:
        return _check(
            "output_log_correlation",
            MechanicalStatus.PASS,
            "Output/log correlation is not applicable to an unaccepted run.",
        )
    if path is None or not mapped_readable:
        return _check(
            "output_log_correlation",
            MechanicalStatus.FAIL,
            "Output/log correlation could not run without a readable mapped workbook.",
        )
    if not isinstance(run_log, Mapping):
        return _check(
            "output_log_correlation",
            MechanicalStatus.FAIL,
            "Output/log correlation could not run without a readable run log.",
        )

    issues: list[str] = []
    issue_count = 0

    def add_issue(message: str) -> None:
        nonlocal issue_count
        issue_count += 1
        if len(issues) < _DETAIL_LIMIT:
            issues.append(message)

    try:
        # The COA check uses random cell access. A normal workbook load is much
        # faster here than repeatedly seeking through a read-only worksheet.
        book = openpyxl.load_workbook(path, read_only=False, data_only=False)
        try:
            coa_sheet = book["COA"]
            notes_sheet = book["Run Notes"]
            expected_ids = canonical_coa_ids()
            actual_ids = [
                str(
                    coa_sheet.cell(
                        row=FIRST_ACCOUNT_ROW + offset,
                        column=ID_COL,
                    ).value
                    or ""
                )
                for offset in range(len(expected_ids))
            ]
            if actual_ids != expected_ids:
                first = next(
                    (
                        position
                        for position, (actual, expected) in enumerate(
                            zip(actual_ids, expected_ids)
                        )
                        if actual != expected
                    ),
                    None,
                )
                detail = (
                    f" at output row {FIRST_ACCOUNT_ROW + first}"
                    if first is not None
                    else ""
                )
                add_issue(f"COA account order differs from the current chart{detail}")
            row_by_id = {
                coa_id: FIRST_ACCOUNT_ROW + position
                for position, coa_id in enumerate(actual_ids)
                if coa_id
            }

            source = run_log.get("source")
            source = source if isinstance(source, Mapping) else {}
            expected_source_name = source.get("name")
            actual_source_name = notes_sheet.cell(row=4, column=3).value
            if actual_source_name != expected_source_name:
                add_issue(
                    "Run Notes source name differs: "
                    f"workbook={actual_source_name!r}, log={expected_source_name!r}"
                )

            periods = source.get("periods")
            periods = periods if _is_sequence(periods) else []
            labels: list[str] = []
            period_ids: list[str] = []
            for period in periods:
                if not isinstance(period, Mapping):
                    continue
                labels.append(str(period.get("label") or ""))
                period_ids.append(str(period.get("period_id") or ""))
            for offset, label in enumerate(labels):
                actual = coa_sheet.cell(
                    row=HEADER_ROW,
                    column=FIRST_PERIOD_COL + offset,
                ).value
                if actual != label:
                    add_issue(
                        f"period {period_ids[offset]!r} header differs: "
                        f"workbook={actual!r}, log={label!r}"
                    )
            actual_period_note = notes_sheet.cell(row=5, column=3).value
            expected_period_note = ", ".join(labels)
            if actual_period_note != expected_period_note:
                add_issue(
                    "Run Notes period text differs: "
                    f"workbook={actual_period_note!r}, log={expected_period_note!r}"
                )

            outcome = run_log.get("outcome")
            outcome = outcome if isinstance(outcome, Mapping) else {}
            expected_mapped_count = outcome.get("accounts_populated")
            actual_mapped_count = notes_sheet.cell(row=7, column=3).value
            if (
                expected_mapped_count is not None
                and actual_mapped_count != expected_mapped_count
            ):
                add_issue(
                    "Run Notes mapped-account count differs: "
                    f"workbook={actual_mapped_count!r}, "
                    f"log={expected_mapped_count!r}"
                )

            values_by_period = run_log.get("values_by_period")
            values_by_period = (
                values_by_period if isinstance(values_by_period, Mapping) else {}
            )
            for offset, period_id in enumerate(period_ids):
                expected_values = values_by_period.get(period_id)
                if not isinstance(expected_values, Mapping):
                    continue
                for coa_id, expected in expected_values.items():
                    row = row_by_id.get(str(coa_id))
                    if row is None:
                        add_issue(f"log account {coa_id!r} is absent from output COA")
                        continue
                    target = coa_sheet.cell(
                        row=row,
                        column=FIRST_PERIOD_COL + offset,
                    )
                    actual = target.value
                    if not _same_written_value(
                        actual,
                        expected,
                        coa_id=str(coa_id),
                        number_format=target.number_format,
                    ):
                        add_issue(
                            f"{period_id}/{coa_id} differs: "
                            f"workbook={actual!r}, log={expected!r}"
                        )
        finally:
            book.close()
    except Exception as exc:
        return _check(
            "output_log_correlation",
            MechanicalStatus.FAIL,
            "Mapped workbook could not be correlated with the run log.",
            [str(path), f"{type(exc).__name__}: {exc}"],
        )

    if issue_count:
        evidence = [str(path), *_limited(issues)]
        omitted = issue_count - len(issues)
        if omitted > 0:
            evidence.append(f"{omitted} additional mismatch(es) omitted")
        return _check(
            "output_log_correlation",
            MechanicalStatus.FAIL,
            f"Mapped workbook differs from the run log in {issue_count} place(s).",
            evidence,
        )
    return _check(
        "output_log_correlation",
        MechanicalStatus.PASS,
        "Mapped workbook metadata and written COA values match the run log.",
        [str(path), f"periods={len(period_ids)}"],
    )


def _validate_supplied_index(
    run_log: Mapping[str, Any],
    index: EvidenceIndex,
) -> None:
    raw_rows = run_log.get("evidence_rows")
    raw_accounts = run_log.get("accounts")
    if not _is_sequence(raw_rows) or not _is_sequence(raw_accounts):
        raise EvidenceIndexError("run log evidence_rows and accounts must be lists")
    expected_rows = tuple(
        str(row.get("row_key") or "")
        for row in raw_rows
        if isinstance(row, Mapping)
    )
    expected_accounts = tuple(
        str(account.get("coa_id") or "")
        for account in raw_accounts
        if isinstance(account, Mapping)
    )
    if expected_rows != tuple(row.row_key for row in index.rows):
        raise EvidenceIndexError("supplied EvidenceIndex does not match evidence_rows")
    if expected_accounts != tuple(account.coa_id for account in index.accounts):
        raise EvidenceIndexError("supplied EvidenceIndex does not match accounts")


def _same_written_value(
    actual: Any,
    expected: Any,
    *,
    coa_id: str,
    number_format: str,
) -> bool:
    if expected is None:
        return actual in (None, "")
    if actual is None or isinstance(actual, bool) or isinstance(expected, bool):
        return False
    try:
        expected_written = _output_number(
            coa_id,
            expected,
            number_format=number_format,
        )
        actual_number = float(actual)
    except (TypeError, ValueError, OverflowError):
        return False
    return math.isfinite(expected_written) and math.isclose(
        actual_number,
        expected_written,
        rel_tol=0.0,
        abs_tol=0.000001,
    )


def _mapping_accepted(run_log: Mapping[str, Any] | None) -> bool | None:
    if not isinstance(run_log, Mapping):
        return None
    outcome = run_log.get("outcome")
    if not isinstance(outcome, Mapping):
        return None
    accepted = outcome.get("accepted")
    return accepted if isinstance(accepted, bool) else None


def _limited(items: Sequence[str]) -> list[str]:
    values = list(items)
    if len(values) <= _DETAIL_LIMIT:
        return values
    return [*values[:_DETAIL_LIMIT], f"{len(values) - _DETAIL_LIMIT} more omitted"]


def _is_sequence(value: Any) -> bool:
    return isinstance(value, Sequence) and not isinstance(
        value,
        (str, bytes, bytearray),
    )


def _check(
    code: str,
    status: MechanicalStatus,
    message: str,
    evidence: Sequence[str] | None = None,
) -> MechanicalCheck:
    return MechanicalCheck(
        code=code,
        status=status,
        message=message,
        evidence=list(evidence or []),
    )


__all__ = ["run_mechanical_checks"]
