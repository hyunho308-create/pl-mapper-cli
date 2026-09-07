from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

import openpyxl
import pytest
from pypdf import PdfWriter

from hotel_pl_normalizer.evaluation.evidence import (
    EvidenceIndexError,
    build_evidence_index,
)
from hotel_pl_normalizer.evaluation.mechanical import run_mechanical_checks
from hotel_pl_normalizer.evaluation.models import MechanicalStatus
from hotel_pl_normalizer.mapping.coa import canonical_coa_ids
from hotel_pl_normalizer.output import (
    FIRST_ACCOUNT_ROW,
    FIRST_PERIOD_COL,
    HEADER_ROW,
    ID_COL,
)

TARGET = "S1.transient_rooms_revenue"


def _excel_row(row: int, label: str, value: float) -> dict:
    row_key = f"P&L!{row}"
    return {
        "row_key": row_key,
        "label": label,
        "selected_value_columns": {"actual": 3},
        "selected_values": {"actual": value},
        "selected_value_column": 3,
        "selected_value": value,
        "selected_value_formats": {"actual": "$#,##0"},
        "indent": 0,
        "bold": False,
        "label_column": 1,
        "label_rule": "leftmost_text",
        "label_status": "found",
        "label_context": [],
        "locator": {
            "kind": "excel",
            "identity": f"excel:{row_key}",
            "display": f"P&L row {row}",
            "sheet_name": "P&L",
            "row_index": row,
        },
        "anchors_by_period": {
            "actual": {
                "kind": "excel",
                "column_index": 3,
                "excel_column": "C",
                "display": "C",
            }
        },
    }


def _pdf_row(page: int, line: int, label: str, value: float) -> dict:
    row_key = f"Page {page:03d}!{line}"
    return {
        "row_key": row_key,
        "label": label,
        "selected_value_columns": {"actual": "x=410.000"},
        "selected_values": {"actual": value},
        "selected_value_column": "x=410.000",
        "selected_value": value,
        "indent": 0,
        "bold": False,
        "label_x0": 42.0,
        "label_rule": "pdf_text",
        "label_status": "found",
        "label_context": [],
        "pdf_source": {"page": page, "line_id": f"p{page}-l{line}", "top": 80.0},
        "locator": {
            "kind": "pdf",
            "identity": f"pdf:page={page}:line={line}",
            "display": f"Page {page}, line {line}",
            "page_number": page,
            "line_number": line,
            "line_id": f"p{page}-l{line}",
            "top": 80.0,
        },
        "anchors_by_period": {
            "actual": {
                "kind": "pdf",
                "right_edge": 410.0,
                "display": "x=410.000",
            }
        },
    }


def _row_ref(row: dict, *, found: bool = True) -> dict:
    return {
        "row_key": row["row_key"],
        "label": row["label"] if found else None,
        "value": row["selected_value"] if found else None,
        "indent": row.get("indent") if found else None,
        "bold": row.get("bold") if found else None,
        "found": found,
        **(
            {
                "locator": row["locator"],
                "anchors_by_period": row["anchors_by_period"],
            }
            if found
            else {}
        ),
    }


def _run_log(source_name: str, *, value: float = 125.5) -> dict:
    evidence = [
        _excel_row(4, "Rooms Revenue", value),
        _excel_row(5, "Transient Rooms Revenue", value),
        _excel_row(6, "Group Rooms Revenue", 0),
    ]
    return {
        "log_version": 5,
        "source": {
            "name": source_name,
            "workbook_id": "wb-test",
            "period": "2026 Actual",
            "periods": [{"period_id": "actual", "label": "2026 Actual"}],
        },
        "outcome": {
            "accepted": True,
            "classification": "clean",
            "accounts_populated": 1,
        },
        "evidence_rows": evidence,
        "accounts": [
            {
                "coa_id": TARGET,
                "account_name": "Transient Rooms Revenue",
                "department": "Rooms",
                "computed_value": value,
                "computed_values": {"actual": value},
                "operation": "direct",
                "source_rows": [_row_ref(evidence[1])],
                "excluded_rows": [],
            }
        ],
        "values_by_period": {"actual": {TARGET: value}},
        "feedback_manifest": {
            "findings": [],
            "inputs": [],
            "rendered_count": 0,
            "unmatched_count": 0,
        },
    }


def _write_source(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "P&L"
    sheet["A5"] = "Transient Rooms Revenue"
    sheet["C5"] = 125.5
    book.save(path)
    book.close()


def _write_mapped(path: Path, run_log: dict, *, value: float | None = None) -> None:
    book = openpyxl.Workbook()
    coa = book.active
    coa.title = "COA"
    coa["B2"] = "COA_ID"
    coa["W2"] = "Mapped Labels"
    coa["X2"] = "MODEL FEEDBACK"
    book.create_sheet("KHP Model Accounts")
    notes = book.create_sheet("Run Notes")
    ids = canonical_coa_ids()
    for offset, coa_id in enumerate(ids):
        coa.cell(row=FIRST_ACCOUNT_ROW + offset, column=ID_COL, value=coa_id)
    periods = run_log["source"]["periods"]
    for offset, period in enumerate(periods):
        coa.cell(
            row=HEADER_ROW,
            column=FIRST_PERIOD_COL + offset,
            value=period["label"],
        )
        period_values = run_log["values_by_period"][period["period_id"]]
        for coa_id, expected in period_values.items():
            row = FIRST_ACCOUNT_ROW + ids.index(coa_id)
            written = expected if value is None else value
            coa.cell(row=row, column=FIRST_PERIOD_COL + offset, value=written)
    notes["C4"] = run_log["source"]["name"]
    notes["C5"] = ", ".join(period["label"] for period in periods)
    notes["C7"] = run_log["outcome"]["accounts_populated"]
    book.save(path)
    book.close()


def _write_log(path: Path, run_log: dict) -> None:
    path.write_text(json.dumps(run_log), encoding="utf-8")


def _by_code(checks) -> dict:
    return {check.code: check for check in checks}


def test_evidence_index_resolves_exact_rows_and_same_scope_context() -> None:
    log = _run_log("source.xlsx")
    excluded = _excel_row(9, "Less: Package Revenue", 10)
    other_sheet = deepcopy(_excel_row(2, "Unrelated", 1))
    other_sheet["row_key"] = "Detail!2"
    other_sheet["locator"].update(
        {
            "identity": "excel:Detail!2",
            "display": "Detail row 2",
            "sheet_name": "Detail",
            "row_index": 2,
        }
    )
    log["evidence_rows"].extend([excluded, other_sheet])
    log["accounts"][0]["excluded_rows"] = [_row_ref(excluded)]

    index = build_evidence_index(log)

    assert [row.row_key for row in index.rows_for_account(TARGET)] == ["P&L!5"]
    assert index.account(TARGET).excluded_row_keys == ("P&L!9",)
    assert [row.row_key for row in index.nearby("P&L!5", radius=1)] == [
        "P&L!4",
        "P&L!5",
        "P&L!6",
    ]
    assert [row.row_key for row in index.context_for_account(TARGET, radius=1)] == [
        "P&L!4",
        "P&L!5",
        "P&L!6",
        "P&L!9",
    ]
    assert index.missing_row_keys_for_account(TARGET) == ()


def test_evidence_index_keeps_pdf_context_on_the_same_page() -> None:
    rows = [
        _pdf_row(2, 1, "Rooms", 100),
        _pdf_row(2, 2, "Transient", 100),
        _pdf_row(3, 1, "Food", 50),
    ]
    log = _run_log("source.pdf")
    log["evidence_rows"] = rows
    log["accounts"][0]["source_rows"] = [_row_ref(rows[1])]

    index = build_evidence_index(log)

    assert [row.row_key for row in index.nearby("Page 002!2", radius=2)] == [
        "Page 002!1",
        "Page 002!2",
    ]
    assert index.row("Page 003!1").locator.page_number == 3


def test_evidence_index_rejects_disagreement_between_key_and_locator() -> None:
    log = _run_log("source.xlsx")
    log["evidence_rows"][0]["locator"]["row_index"] = 99

    with pytest.raises(EvidenceIndexError, match="disagrees with its typed locator"):
        build_evidence_index(log)


def test_mechanical_checks_pass_for_a_correlated_single_pl_run(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    run_log_path = tmp_path / "run_log.json"
    mapped = tmp_path / "source [MAPPED].xlsx"
    log = _run_log(source.name)
    _write_source(source)
    _write_log(run_log_path, log)
    _write_mapped(mapped, log)

    checks = _by_code(run_mechanical_checks(source, run_log_path, mapped))

    assert set(checks) == {
        "source_artifact",
        "run_log_artifact",
        "run_log_integrity",
        "mapped_workbook_artifact",
        "output_log_correlation",
    }
    assert all(check.status == MechanicalStatus.PASS for check in checks.values())


def test_mechanical_checks_preserve_ratio_precision(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    run_log_path = tmp_path / "run_log.json"
    mapped = tmp_path / "source [MAPPED].xlsx"
    log = _run_log(source.name)
    log["values_by_period"]["actual"]["S12.occupancy"] = 3.2437
    _write_source(source)
    _write_log(run_log_path, log)
    _write_mapped(mapped, log)

    checks = _by_code(run_mechanical_checks(source, run_log_path, mapped))

    assert checks["output_log_correlation"].status == MechanicalStatus.PASS


def test_mechanical_checks_fail_closed_for_an_old_run_log(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    run_log_path = tmp_path / "run_log.json"
    mapped = tmp_path / "source [MAPPED].xlsx"
    log = _run_log(source.name)
    log["log_version"] = 4
    _write_source(source)
    _write_log(run_log_path, log)
    _write_mapped(mapped, log)

    checks = _by_code(run_mechanical_checks(source, run_log_path, mapped))

    assert checks["run_log_artifact"].status == MechanicalStatus.PASS
    assert checks["run_log_integrity"].status == MechanicalStatus.FAIL
    assert "version 5" in " ".join(checks["run_log_integrity"].evidence)


def test_mechanical_checks_detect_post_write_value_changes(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    run_log_path = tmp_path / "run_log.json"
    mapped = tmp_path / "source [MAPPED].xlsx"
    log = _run_log(source.name)
    _write_source(source)
    _write_log(run_log_path, log)
    _write_mapped(mapped, log, value=999)

    checks = _by_code(run_mechanical_checks(source, run_log_path, mapped))

    correlation = checks["output_log_correlation"]
    assert correlation.status == MechanicalStatus.FAIL
    assert any(TARGET in item for item in correlation.evidence)


def test_missing_account_evidence_does_not_block_direct_file_review(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    run_log_path = tmp_path / "run_log.json"
    mapped = tmp_path / "source [MAPPED].xlsx"
    log = _run_log(source.name)
    missing = _excel_row(99, "Missing Row", 125.5)
    log["accounts"][0]["source_rows"] = [_row_ref(missing, found=False)]
    _write_source(source)
    _write_log(run_log_path, log)
    _write_mapped(mapped, log)

    checks = _by_code(run_mechanical_checks(source, run_log_path, mapped))

    assert checks["run_log_integrity"].status == MechanicalStatus.PASS
    assert "judge_evidence_completeness" not in checks
    assert all(check.status == MechanicalStatus.PASS for check in checks.values())


def test_pdf_source_readability_uses_the_same_per_pl_checks(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with source.open("wb") as stream:
        writer.write(stream)
    run_log_path = tmp_path / "run_log.json"
    mapped = tmp_path / "source [MAPPED].xlsx"
    log = _run_log(source.name)
    pdf_row = _pdf_row(1, 1, "Transient Rooms Revenue", 125.5)
    log["evidence_rows"] = [pdf_row]
    log["accounts"][0]["source_rows"] = [_row_ref(pdf_row)]
    _write_log(run_log_path, log)
    _write_mapped(mapped, log)

    checks = _by_code(run_mechanical_checks(source, run_log_path, mapped))

    assert checks["source_artifact"].status == MechanicalStatus.PASS
    assert "judge_evidence_completeness" not in checks
