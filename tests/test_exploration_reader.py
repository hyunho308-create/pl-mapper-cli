from __future__ import annotations

import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

from hotel_pl_normalizer.structure.exploration.reader import LazyWorkbook
from hotel_pl_normalizer.structure.period_headers import latest_header_month


def _reverse_workbook_sheet_order(path: Path) -> None:
    """Reorder workbook tabs while retaining their relationship targets."""
    staged = path.with_suffix(".reordered.xlsx")
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(staged, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                root = ET.fromstring(payload)
                sheets = root.find(f"{{{namespace}}}sheets")
                assert sheets is not None
                children = list(sheets)
                sheets[:] = reversed(children)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, payload)
    staged.replace(path)


def test_merged_ranges_follow_sheet_relationship_after_tab_reordering(tmp_path: Path):
    path = tmp_path / "reordered.xlsx"
    workbook = openpyxl.Workbook()
    alpha = workbook.active
    alpha.title = "Alpha"
    alpha["A1"] = "Alpha header"
    alpha.merge_cells("A1:B1")
    beta = workbook.create_sheet("Beta")
    beta["C1"] = "Beta header"
    beta.merge_cells("C1:D1")
    workbook.save(path)
    _reverse_workbook_sheet_order(path)

    with LazyWorkbook(path) as reader:
        assert [sheet.sheet_name for sheet in reader.sheets()] == ["Beta", "Alpha"]
        assert reader.merged_ranges("Beta") == ["C1:D1"]
        assert reader.merged_ranges("Alpha") == ["A1:B1"]


def test_latest_header_month_reads_full_textual_dates_without_using_day_as_year():
    assert latest_header_month(["For the Month Ending January 31, 2025"]) == (
        2025,
        1,
    )
    assert latest_header_month(["For the Month Ending January 31 2025"]) == (
        2025,
        1,
    )
