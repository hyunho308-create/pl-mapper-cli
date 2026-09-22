"""Presentation cleanup must not mutate accounting evidence or model formulas."""
from copy import deepcopy

from openpyxl import load_workbook

from hotel_pl_normalizer.mapping.coa import load_coa
from hotel_pl_normalizer.mapping.findings import Finding
from hotel_pl_normalizer.output import _incomplete_detail_omissions, write_normalized_workbook
from hotel_pl_normalizer.pipeline import NormalizationResult


def result():
    return NormalizationResult(
        workbook_id="test", source_name="test.xlsx", period_label="Actual", values={},
        coa=load_coa(), accepted=True,
        period_labels={"a": "2025 Actual", "b": "2025 Budget"},
        period_values={"a": {"S1.nonmanagement": 100, "S1.front_office": 70},
                       "b": {"S1.nonmanagement": 100, "S1.front_office": 100}},
        checks_by_period={"a": [Finding("warning", "source_detail_incomplete", "S1.nonmanagement",
                                       {"parent":100,"children":70,"variance":30})]},
        run_summary="A long summary of source differences. " * 12,
    )


def test_period_specific_cleanup_preserves_audit_and_model(tmp_path):
    r = result()
    original = result()
    baseline = result(); baseline.checks_by_period = {}
    before = load_workbook(write_normalized_workbook(baseline, tmp_path / "before.xlsx"))
    after = load_workbook(write_normalized_workbook(r, tmp_path / "after.xlsx"))
    rows = {row[1].value:row[1].row for row in after["COA"].iter_rows(min_row=3)}
    child, parent = rows["S1.front_office"], rows["S1.nonmanagement"]
    assert after["COA"].cell(child,3).value is None
    assert after["COA"].cell(child,4).value == 100
    assert after["COA"].cell(child,3).fill.fgColor.rgb != "00FFFF00"
    assert after["COA"].cell(parent,3).value == 100
    assert after["COA"].cell(parent,24).value == "Detailed breakdown omitted because it does not add up to the total."
    assert after["COA"].cell(child,24).value is None
    assert after["Run Notes"]["C9"].alignment.wrap_text
    assert after["Run Notes"].row_dimensions[9].height >= 100
    for row in before["KHP Model Accounts"]:
        for cell in row:
            other = after["KHP Model Accounts"][cell.coordinate]
            assert other.value == cell.value and other._style == cell._style
    assert r.period_values == original.period_values
    assert r.checks_by_period == original.checks_by_period
    assert r.accepted == original.accepted
    before.close(); after.close()


def test_residual_and_tolerance_preserve_children():
    r=result()
    r.coa["S1.housekeeping"]["is_residual"]="true"
    assert _incomplete_detail_omissions(r,[("a","Actual",r.period_values["a"])]) == (set(),set())
    r.coa["S1.housekeeping"]["is_residual"]="false"
    r.period_values["a"]["S1.front_office"]=96
    assert _incomplete_detail_omissions(r,[("a","Actual",r.period_values["a"])]) == (set(),set())


def test_department_subtotals_anchor_independent_detail():
    r=result()
    total="S2.total_food_and_beverage_expenses"
    wages="S2.salaries_and_wages"
    r.period_values={"a":{total:1000,"S2.labor_costs_and_related_expenses":500,
        "S2.other_expenses":300,"S2.total_cost_of_sales_and_other_revenue":150,
        wages:400,"S2.management":100,"S2.service_management":100}}
    r.checks_by_period={"a":[Finding("warning","source_detail_incomplete",p) for p in (total,wages)]}
    omitted, collapsed=_incomplete_detail_omissions(r,[("a","Actual",r.period_values["a"])])
    assert (wages,"a") in collapsed
    assert ("S2.management","a") in omitted
    assert ("S2.service_management","a") in omitted
    assert not any((key,"a") in omitted for key in
                   (total,wages,"S2.labor_costs_and_related_expenses","S2.other_expenses","S2.total_cost_of_sales_and_other_revenue"))
