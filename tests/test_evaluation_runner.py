from __future__ import annotations

import json
import subprocess
from copy import deepcopy

import openpyxl
import pytest
from test_evaluation_mechanical import _excel_row, _row_ref, _write_mapped

from hotel_pl_normalizer.evaluation.grouping import group_findings
from hotel_pl_normalizer.evaluation.report import render_evaluation_markdown
from hotel_pl_normalizer.evaluation.runner import evaluate_run
from hotel_pl_normalizer.mapping.coa import load_coa
from hotel_pl_normalizer.mapping.findings import Finding

PARENT = "S1.other_expenses"
LAUNDRY = "S1.laundry_and_dry_cleaning"
UNIFORMS = "S1.rooms_all_other"


@pytest.fixture
def review_log():
    rows = [
        _excel_row(4, "Total Other Expenses", 100),
        _excel_row(5, "Laundry", 70),
        _excel_row(6, "Uniform laundry", 30),
    ]
    for row in rows:
        row["selected_values"]["prior"] = row["selected_value"]
        row["selected_value_columns"]["prior"] = 4
        row["anchors_by_period"]["prior"] = {
            "kind": "excel",
            "column_index": 4,
            "excel_column": "D",
            "display": "D",
        }
    values = {PARENT: 100, LAUNDRY: 70, UNIFORMS: 30}
    return {
        "log_version": 5,
        "source": {
            "name": "source.xlsx",
            "workbook_id": "wb-test",
            "periods": [
                {"period_id": "actual", "label": "Actual"},
                {"period_id": "prior", "label": "Prior"},
            ],
        },
        "outcome": {"accepted": True, "accounts_populated": 3},
        "evidence_rows": rows,
        "accounts": [
            {
                "coa_id": cid,
                "operation": "direct",
                "source_rows": [_row_ref(row)],
                "excluded_rows": [],
                "computed_values": {"actual": value, "prior": value},
            }
            for (cid, value), row in zip(values.items(), rows)
        ],
        "values_by_period": {"actual": values, "prior": dict(values)},
        "feedback_manifest": {"findings": []},
    }


@pytest.fixture
def run_review(tmp_path):
    def run(log, *, expected=None):
        source, mapped, log_path = (
            tmp_path / name for name in ("source.xlsx", "mapped.xlsx", "run_log.json")
        )
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "P&L"
        for row in log["evidence_rows"]:
            number = row["locator"]["row_index"]
            sheet.cell(number, 1, row["label"])
            for pid, column in row["selected_value_columns"].items():
                sheet.cell(number, column, row["selected_values"].get(pid))
        book.save(source)
        book.close()
        _write_mapped(mapped, log)
        log_path.write_text(json.dumps(log), encoding="utf-8")
        before = {p: p.read_bytes() for p in (source, mapped, log_path)}
        result = evaluate_run(source, log_path, mapped, expected_period_ids=expected)
        assert {p: p.read_bytes() for p in before} == before
        assert set(tmp_path.iterdir()) == set(before)
        return result

    return run


def test_review_is_local_and_never_approves(review_log, run_review, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Quick evaluation must not launch any subprocess")

    monkeypatch.setattr(subprocess, "run", forbidden)
    monkeypatch.setattr(subprocess, "Popen", forbidden)
    result = run_review(review_log)
    assert result.status == "ready_for_review"
    assert {c["coa_id"] for c in result.children} == {LAUNDRY, UNIFORMS}
    text = render_evaluation_markdown(result)
    assert "Semantic review: **pending**" in text
    assert "END OF REVIEW EVIDENCE" in text
    assert "approved" not in text


@pytest.mark.parametrize("record_drop", [True, False])
def test_missing_requested_period(review_log, run_review, record_drop):
    review_log["source"]["periods"].pop()
    review_log["values_by_period"].pop("prior")
    if record_drop:
        review_log["dropped_periods"] = {
            "prior": "Could not bind the prior-year column"
        }
    result = run_review(review_log, expected=[] if record_drop else ["actual", "prior"])
    assert result.status == "issues_found"
    assert any(
        f["code"] == "missing_period" and f["periods"] == ["prior"]
        for f in result.findings
    )


def test_child_period_blank_despite_source_amount(review_log, run_review):
    review_log["values_by_period"]["prior"][LAUNDRY] = None
    result = run_review(review_log)
    finding = next(f for f in result.findings if f["code"] == "period_value_missing")
    assert finding["targets"] == [LAUNDRY]
    assert finding["periods"] == ["prior"]
    assert finding["refs"] == ["P&L!5"]


def test_same_label_period_alternative_stays_visible(review_log, run_review):
    review_log["values_by_period"]["prior"][LAUNDRY] = None
    review_log["evidence_rows"][1]["selected_values"]["prior"] = None
    alternative = deepcopy(review_log["evidence_rows"][1])
    alternative["row_key"] = "P&L!8"
    alternative["locator"].update(
        identity="excel:P&L!8", display="P&L row 8", row_index=8
    )
    alternative["selected_values"] = {"actual": None, "prior": 70}
    review_log["evidence_rows"].append(alternative)
    result = run_review(review_log)
    assert any(
        f["code"] == "period_detail_candidate" and "P&L!8" in f["refs"]
        for f in result.findings
    )


def test_dropped_detail_used_only_by_parent_is_visible(review_log, run_review):
    review_log["accounts"][0]["source_rows"].append(
        _row_ref(review_log["evidence_rows"][2])
    )
    review_log["accounts"][2].update(operation="no_value", source_rows=[])
    for values in review_log["values_by_period"].values():
        values[UNIFORMS] = None
    result = run_review(review_log)
    row = next(r for r in result.rows if r["key"] == "P&L!6")
    assert row["candidate"] and row["usage"] == "parent"
    assert row["values"] == [30, 30]


def test_summary_only_source_is_a_question_not_a_defect(review_log, run_review):
    review_log["evidence_rows"] = review_log["evidence_rows"][:1]
    review_log["accounts"] = review_log["accounts"][:1]
    review_log["accounts"][0]["child_coverage"] = "not_present"
    for values in review_log["values_by_period"].values():
        values[LAUNDRY] = values[UNIFORMS] = None
    result = run_review(review_log)
    assert result.status == "ready_for_review"
    assert all(f["severity"] == "warning" for f in result.findings)
    assert any(f["code"] == "parent_without_detail" for f in result.findings)


def test_sibling_swap_cannot_hide_behind_matching_total(review_log, run_review):
    first, second = review_log["accounts"][1:]
    first["source_rows"], second["source_rows"] = (
        second["source_rows"],
        first["source_rows"],
    )
    for values in review_log["values_by_period"].values():
        values[LAUNDRY], values[UNIFORMS] = values[UNIFORMS], values[LAUNDRY]
    result = run_review(review_log)
    child = next(c for c in result.children if c["coa_id"] == LAUNDRY)
    assert child["refs"] == ["P&L!6"]
    assert (
        next(r for r in result.rows if r["key"] == child["refs"][0])["label"]
        == "Uniform laundry"
    )
    assert {LAUNDRY, UNIFORMS} <= {d["coa_id"] for d in result.definitions}
    assert "EVERY active child" in render_evaluation_markdown(result)


def test_zero_is_distinct_from_missing_and_unbacked_child_is_included(
    review_log, run_review
):
    review_log["values_by_period"]["prior"][LAUNDRY] = 0
    review_log["accounts"][2]["source_rows"] = []
    result = run_review(review_log)
    assert any(c["coa_id"] == UNIFORMS for c in result.children)
    assert any(f["code"] == "zero_from_nonzero_source" for f in result.findings)
    assert any(f["code"] == "child_without_rows" for f in result.findings)


def test_raw_error_survives_missing_feedback_entry(review_log, run_review):
    error = Finding(
        "error", "hierarchy_complete", PARENT, {"variance": 30}, period_id="prior"
    )
    review_log["findings"] = [error.to_dict(), error.to_dict()]
    review_log["findings_by_period"] = {"prior": [error.to_dict()]}
    review_log["checks_by_period"] = {"prior": [str(error)]}
    review_log["execution_issues"] = [
        "Failed to calculate period",
        "Failed to calculate period",
    ]
    result = run_review(review_log)
    assert result.status == "issues_found"
    assert len([f for f in result.findings if f["code"] == "hierarchy_complete"]) == 1
    assert len([f for f in result.findings if f["code"] == "execution_error"]) == 1


def test_warning_resolves_rows_even_when_not_used_by_children(review_log, run_review):
    review_log["feedback_manifest"]["findings"] = [
        {
            "finding_id": "source-warning",
            "severity": "warning",
            "explanation": "Source discrepancy",
            "source_refs": ["P&L!4"],
            "affected_coa_ids": [PARENT],
            "periods": [],
        }
    ]
    result = run_review(review_log)
    assert (
        "P&L!4"
        in next(f for f in result.findings if f["code"] == "source-warning")["refs"]
    )
    assert any(r["key"] == "P&L!4" and r["values"] == [100, 100] for r in result.rows)


def test_missing_warning_reference_is_incomplete(review_log, run_review):
    review_log["feedback_manifest"]["findings"] = [
        {
            "finding_id": "missing",
            "severity": "warning",
            "source_refs": ["P&L!999"],
        }
    ]
    result = run_review(review_log)
    assert result.status == "incomplete"
    assert "MISSING:P&L!999" in render_evaluation_markdown(result)


def test_unreadable_log_returns_incomplete(tmp_path):
    path = tmp_path / "run_log.json"
    path.write_text("[]", encoding="utf-8")
    result = evaluate_run(tmp_path / "source.xlsx", path, None)
    assert result.status == "incomplete"
    assert "Cannot prepare evidence" in render_evaluation_markdown(result)


def test_missing_parent_reference_is_incomplete(review_log, run_review):
    review_log["accounts"][0]["source_rows"] = [{"row_key": "P&L!999"}]
    result = run_review(review_log)
    assert result.status == "incomplete"
    assert any(f["code"] == "missing_source_references" for f in result.findings)


def test_parent_subtraction_does_not_hide_available_detail(review_log, run_review):
    review_log["accounts"][2].update(operation="no_value", source_rows=[])
    review_log["accounts"][0]["excluded_rows"] = [_row_ref(review_log["evidence_rows"][2])]
    for values in review_log["values_by_period"].values():
        values[UNIFORMS] = None
    result = run_review(review_log)
    row = next(r for r in result.rows if r["key"] == "P&L!6")
    assert row["candidate"] and row["usage"] == "parent subtraction"


def test_absent_feedback_manifest_still_exposes_raw_errors(review_log, run_review):
    review_log.pop("feedback_manifest")
    review_log["execution_issues"] = ["Mapping calculation failed"]
    result = run_review(review_log)
    assert result.status == "incomplete"
    assert any(f["code"] == "execution_error" for f in result.findings)


def test_review_includes_authoritative_synonyms(review_log, run_review):
    result = run_review(review_log)
    definition = next(item for item in result.definitions if item["coa_id"] == LAUNDRY)
    assert definition["synonyms"] == load_coa()[LAUNDRY]["synonyms"]
    assert definition["synonyms"].replace("|", "&#124;") in render_evaluation_markdown(result)


def test_feedback_and_raw_copies_group_without_losing_period_amounts(review_log, run_review):
    review_log["feedback_manifest"]["findings"] = [{
        "finding_id": "source-conflict", "severity": "warning",
        "primary_coa_id": PARENT, "affected_coa_ids": [PARENT],
        "source_refs": ["P&L!4", "P&L!5"], "explanation": "Preserve both source layers.",
        "periods": [
            {"period_id": "actual", "selected_value": 100, "comparison_value": 90, "variance": 10},
            {"period_id": "prior", "selected_value": 100, "comparison_value": 80, "variance": 20},
        ],
    }]
    checks = [Finding("warning", "hierarchy_complete", PARENT,
                      {"parent": 100, "children": amount, "variance": 100 - amount}, period_id=pid)
              for pid, amount in (("actual", 90), ("prior", 80))]
    review_log["findings"] = [item.to_dict() for item in checks]
    review_log["findings_by_period"] = {item.period_id: [item.to_dict()] for item in checks}
    review_log["checks_by_period"] = {item.period_id: [str(item)] for item in checks}
    review_log["checks"] = [f"{item}|period={label}" for item, label in zip(checks, ("Actual", "Prior"))]
    result = run_review(review_log)
    assert len(result.findings) == 1
    issue = result.findings[0]
    assert issue["occurrences"] == 9
    assert issue["periods"] == ["actual", "prior"]
    assert {"feedback", "validator", "check"} == set(issue["origins"])
    assert {"P&L!4", "P&L!5"} <= set(issue["refs"])
    report = render_evaluation_markdown(result)
    assert '"variance":10' in report and '"variance":20' in report
    assert "9 recorded occurrences" in report


def test_equal_amount_conflicts_with_different_source_scopes_stay_separate(review_log, run_review):
    review_log["review_items"] = [
        {"review_item_id": "first", "source_rows": ["P&L!4", "P&L!5"]},
        {"review_item_id": "second", "source_rows": ["P&L!4", "P&L!6"]},
    ]
    review_log["findings"] = [
        Finding("warning", "source_layer_conflict", PARENT,
                {"actual": 100, "expected": 90, "variance": 10},
                period_id="actual", review_item_id=identifier).to_dict()
        for identifier in ("first", "second")
    ]
    result = run_review(review_log)
    assert len(result.findings) == 2
    assert {tuple(item["refs"]) for item in result.findings} == {
        ("P&L!4", "P&L!5"), ("P&L!4", "P&L!6")
    }


def test_repeated_findings_keep_highest_severity(review_log, run_review):
    review_log["findings"] = [
        Finding(severity, "hierarchy_complete", PARENT, {"variance": 30},
                period_id=pid).to_dict()
        for severity, pid in (("warning", "actual"), ("error", "prior"))
    ]
    result = run_review(review_log)
    assert len(result.findings) == 1
    assert result.findings[0]["severity"] == "error"
    assert result.findings[0]["occurrences"] == 2
    assert result.status == "issues_found"
    report = render_evaluation_markdown(result)
    parent_id = next(f"A{i}" for i, item in enumerate(result.definitions, 1) if item["coa_id"] == PARENT)
    assert f"warning / validator / hierarchy_complete|{parent_id},P1|" in report
    assert f"error / validator / hierarchy_complete|{parent_id},P2|" in report


def test_repeated_source_content_keeps_each_location_columns_and_use(review_log, run_review):
    alternative = deepcopy(review_log["evidence_rows"][1])
    alternative["row_key"] = "P&L!8"
    alternative["locator"].update(identity="excel:P&L!8", display="P&L row 8", row_index=8)
    alternative["selected_value_columns"]["prior"] = 5
    alternative["anchors_by_period"]["prior"].update(column_index=5, excel_column="E", display="E")
    review_log["evidence_rows"].append(alternative)
    result = run_review(review_log)
    report = render_evaluation_markdown(result)
    assert len(result.rows) == 4
    assert "4 source rows (3 distinct content rows)" in report
    assert "R2=D1!5, R4=D1!8" in report
    assert "C / D; C / E" in report
    assert "child; unused *" in report
    assert len(result.children) == 2
    assert "|R2|" in report


def test_equal_labels_with_different_periods_or_context_are_not_compressed(review_log, run_review):
    alternative = deepcopy(review_log["evidence_rows"][1])
    alternative["row_key"] = "P&L!8"
    alternative["locator"].update(identity="excel:P&L!8", display="P&L row 8", row_index=8)
    alternative["selected_values"]["prior"] = None
    review_log["evidence_rows"].append(alternative)
    result = run_review(review_log)
    assert "4 source rows (4 distinct content rows)" in render_evaluation_markdown(result)


def test_direct_source_in_overview_is_complete_and_not_repeated(review_log, run_review):
    result = run_review(review_log)
    report = render_evaluation_markdown(result)
    assert report.count("R2=D1!5") == 1
    assert "R2=D1!5 [C / D; child] Laundry" in report
    assert "R3=D1!6 [C / D; child] Uniform laundry" in report
    assert "Source rows" in report and "R1=D1!4" in report


def test_parent_period_overview_preserves_zero(review_log, run_review):
    for cid in (PARENT, LAUNDRY, UNIFORMS):
        review_log["values_by_period"]["prior"][cid] = 0
    result = run_review(review_log)
    parent = next(item for item in result.coverage if item["parent"] == PARENT and item["period"] == "prior")
    assert parent["value"] == 0 and parent["populated"] == 0
    assert "100 / 0|2/" in render_evaluation_markdown(result)


@pytest.mark.parametrize("excluded", [False, True])
def test_shared_source_is_inlined_only_beside_its_matching_direct_child(
    review_log, run_review, excluded
):
    review_log["accounts"][1]["operation"] = "adjusted_subtotal" if excluded else "sum"
    if excluded:
        review_log["accounts"][1]["excluded_rows"] = [_row_ref(review_log["evidence_rows"][2])]
    else:
        review_log["accounts"][1]["source_rows"].append(_row_ref(review_log["evidence_rows"][2]))
    result = run_review(review_log)
    report = render_evaluation_markdown(result)
    direct_line = next(line for line in report.splitlines() if "R3=D1!6" in line)
    assert "30 / 30|direct" in direct_line
    assert report.count("R3=D1!6") == 1


@pytest.mark.parametrize("rule,details", [
    ("occupancy_above_capacity", {"occupancy": 3.24, "rooms_sold": 84154, "rooms_available": 25944}),
    ("invalid_rooms_available", {"rooms_sold": 81146, "rooms_available": 0}),
])
def test_room_statistics_feedback_groups_only_matching_quantities(review_log, run_review, rule, details):
    review_log["feedback_manifest"]["findings"] = [{
        "finding_id": "room-stat", "severity": "warning", "primary_coa_id": PARENT,
        "affected_coa_ids": [PARENT], "explanation": "Review source room statistics.",
        "source_refs": [], "periods": [{"period_id": "actual", **details}],
    }]
    review_log["checks"] = [f"{Finding('warning', rule, PARENT, details)}|period=Actual"]
    result = run_review(review_log)
    assert len(result.findings) == 1 and result.findings[0]["occurrences"] == 2
    assert result.findings[0]["periods"] == ["actual"]
    review_log["feedback_manifest"]["findings"][0]["periods"][0]["rooms_sold"] += 1
    assert len(run_review(review_log).findings) == 2


def test_ambiguous_period_label_is_not_assigned_to_a_period(review_log, run_review):
    for period in review_log["source"]["periods"]:
        period["label"] = "Actual"
    review_log["checks"] = [f"{Finding('warning', 'source_detail_incomplete', PARENT, {'parent': 100, 'children': 90})}|period=Actual"]
    result = run_review(review_log)
    assert result.findings[0]["periods"] == []
    assert result.findings[0]["detail"]["period"] == "Actual"


def test_recorded_period_id_takes_precedence_over_another_period_label(review_log, run_review):
    review_log["source"]["periods"][1]["label"] = "actual"
    review_log["checks"] = [f"{Finding('warning', 'source_detail_incomplete', PARENT, {'parent': 100, 'children': 90})}|period=actual"]
    assert run_review(review_log).findings[0]["periods"] == ["actual"]


def test_group_variants_keep_affected_accounts_and_promote_info_to_warning():
    common = {"origin": "local", "code": "gap", "message": "Review detail.",
              "refs": [], "detail": None, "_identity": ("same",)}
    groups = group_findings([
        {**common, "severity": "info", "targets": [PARENT, LAUNDRY], "periods": ["actual"]},
        {**common, "severity": "warning", "targets": [PARENT, UNIFORMS], "periods": ["prior"]},
    ])
    assert len(groups) == 1 and groups[0]["severity"] == "warning"
    assert groups[0]["occurrences"] == 2
    assert [variant["targets"] for variant in groups[0]["variants"]] == [
        [PARENT, LAUNDRY], [PARENT, UNIFORMS],
    ]


def test_matching_room_amounts_with_distinct_source_scopes_do_not_join(review_log, run_review):
    details = {"occupancy": 3.24, "rooms_sold": 84154, "rooms_available": 25944}
    review_log["feedback_manifest"]["findings"] = [{
        "finding_id": "room-stat", "severity": "warning", "primary_coa_id": PARENT,
        "affected_coa_ids": [PARENT], "explanation": "Review source room statistics.",
        "source_refs": ["P&L!4", "P&L!5"], "periods": [{"period_id": "actual", **details}],
    }]
    raw = Finding("warning", "occupancy_above_capacity", PARENT, details, period_id="actual").to_dict()
    raw["source_refs"] = ["P&L!4", "P&L!6"]
    review_log["findings"] = [raw]
    assert len(run_review(review_log).findings) == 2


def test_source_caption_context_is_visible_without_changing_period_values(review_log, run_review):
    review_log["evidence_rows"][1].update(label="DCUR", label_context=["Laundry", "DCUR"])
    result = run_review(review_log)
    row = next(row for row in result.rows if row["key"] == "P&L!5")
    assert row["label"] == "DCUR / Laundry"
    assert row["values"] == [70, 70]
    assert "DCUR / Laundry" in render_evaluation_markdown(result)
